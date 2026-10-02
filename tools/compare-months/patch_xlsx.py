# -*- coding: utf-8 -*-
# Пересборка выручки по КАНОНУ ВЫГРУЗОК (как дашборд; решение 01.10.2026).
import json
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

S = '/private/tmp/claude-501/-Users-aleksey-pegas-planner/e1757453-d539-4377-b209-1ccd478bae1d/scratchpad'
r2 = json.load(open(S + '/revenue2.json'))
A, Sp = r2['aug'], r2['sep']

path = S + '/Сравнение эффективности август–сентябрь 2026.xlsx'
wb = load_workbook(path)
F = lambda **kw: Font(name='Arial', **kw)
H = F(bold=True, size=11)
GOOD = Font(name='Arial', color='006300', bold=True)
BAD = Font(name='Arial', color='B8352A', bold=True)
HEAD_FILL = PatternFill('solid', fgColor='EFEFEA')
BORDER = Border(bottom=Side(style='thin', color='D8D6CF'))

# ── Сводка: выручечные строки методикой плитки ──
ws = wb['Сводка']
ws['A2'] = ('Выручка — КАНОН ВЫГРУЗОК, как на дашборде: только выгруженные рейсы, дата фактической выгрузки '
            '(фолбэк — расчётная), без НДС по ставкам настроек (обычные 22%, ИП 7%, наличные 0%) · август 31 дн · сентябрь 30 дн')
def dset(row, a, s, fmt='#,##0.0', good_up=True, comment=None):
    dv = None if (a is None or s is None) else s - a
    ws.cell(row=row, column=2, value=a).number_format = fmt
    ws.cell(row=row, column=3, value=s).number_format = fmt
    c = ws.cell(row=row, column=4, value=dv)
    c.number_format = '+' + fmt + ';-' + fmt + ';0'
    c.font = (GOOD if (dv or 0) > 0 == good_up or ((dv or 0) > 0) == good_up else BAD) if dv else F()
    c.font = GOOD if dv is not None and ((dv > 0) == good_up) else (BAD if dv else F())
    if comment is not None: ws.cell(row=row, column=5, value=comment)

dset(5, A['total']['s']/1e6, Sp['total']['s']/1e6, '#,##0.1', True,
     'канон выгрузок; сверяется с плиткой дашборда «выгружено за месяц»')
dset(6, A['total']['s']/1e6/31, Sp['total']['s']/1e6/30, '#,##0.00', True, 'главная строка: темп')
dset(7, A['total']['g']/1e6, Sp['total']['g']/1e6, '#,##0.1', True, 'те же рейсы с НДС — для сверки с 1С')
dset(8, A['total']['n'], Sp['total']['n'], '#,##0', True, 'выгруженных рейсов')
dset(9, A['total']['n']/31, Sp['total']['n']/30, '#,##0.1', True)
dset(10, A['total']['s']/max(1,A['total']['n']), Sp['total']['s']/max(1,Sp['total']['n']), '#,##0', True)

# ── Лист Выручка: перезаписать целиком той же методикой ──
del wb['Выручка']
ws = wb.create_sheet('Выручка', 1)
ws['A1'] = 'Выручка по неделям и клиентам'; ws['A1'].font = F(bold=True, size=14)
ws['A2'] = ('Канон выгрузок, без НДС по ставкам настроек · недели с понедельника · '
            'сентябрь полный: 01–30.09')
ws['A2'].font = F(size=9, color='666666')
def header(row, cols, widths=None):
    for i, name in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=name)
        c.font = H; c.fill = HEAD_FILL; c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical='center')
    if widths:
        for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
header(4, ['Период', 'Неделя (пн)', 'Выручка бНДС, млн', 'Рейсов', 'млн/сутки недели'], [12, 14, 18, 10, 16])
r = 5
for key, label, days in (('aug', 'август', 7), ('sep', 'сентябрь', 7)):
    for w in r2[key]['weeks']:
        ws.cell(row=r, column=1, value=label).font = F()
        ws.cell(row=r, column=2, value=w['wk'][8:] + '.' + w['wk'][5:7]).font = F()
        ws.cell(row=r, column=3, value=w['s']/1e6).number_format = '#,##0.1'
        ws.cell(row=r, column=4, value=w['n']).number_format = '#,##0'
        ws.cell(row=r, column=5, value=w['s']/1e6/7).number_format = '#,##0.2'
        r += 1
r += 1
ws.cell(row=r, column=1, value='Топ клиентов, сравнение (канон выгрузок)').font = H
r += 1
header(r, ['Клиент', 'Август: рейсов', 'млн бНДС', 'Сентябрь: рейсов', 'млн бНДС', 'Δ млн'], [36, 14, 12, 15, 12, 11])
r += 1
augC = {c['name']: c for c in A['clients']}
for c in Sp['clients']:
    a = augC.get(c['name'])
    dv = c['s']/1e6 - (a['s']/1e6 if a else 0)
    vals = [c['name'][:38], a['n'] if a else 0, (a['s']/1e6 if a else 0), c['n'], c['s']/1e6, dv]
    fmts = [None, '#,##0', '#,##0.1', '#,##0', '#,##0.1', '+#,##0.1;-#,##0.1;0']
    for i, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=i, value=v)
        cell.font = F()
        if fmts[i-1]: cell.number_format = fmts[i-1]
    if abs(dv) > 0.05:
        ws.cell(row=r, column=6).font = GOOD if dv > 0 else BAD
    r += 1
wb.save(path)
print('OK')
