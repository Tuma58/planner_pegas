# -*- coding: utf-8 -*-
import json
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

S = '/private/tmp/claude-501/-Users-aleksey-pegas-planner/e1757453-d539-4377-b209-1ccd478bae1d/scratchpad'
d = json.load(open(S + '/halves.json'))
H1, H2 = d['h1'], d['h2']
path = S + '/Сравнение эффективности август–сентябрь 2026.xlsx'
wb = load_workbook(path)
if 'Сентябрь 1-2' in wb.sheetnames: del wb['Сентябрь 1-2']
ws = wb.create_sheet('Сентябрь 1-2', 1)
F = lambda **kw: Font(name='Arial', **kw)
GOOD = Font(name='Arial', color='006300', bold=True)
BAD = Font(name='Arial', color='B8352A', bold=True)
ws['A1'] = 'Сентябрь: первая ↔ вторая половина'; ws['A1'].font = F(bold=True, size=14)
ws['A2'] = 'Окна: 01–13.09 (13 дн) и 14–27.09 (14 дн) включительно · выручка по дате выполнения, без НДС'
ws['A2'].font = F(size=9, color='666666')
cols = ['Показатель', '01–13.09', '14–27.09', 'Δ', 'Комментарий']
for i, name in enumerate(cols, 1):
    c = ws.cell(row=4, column=i, value=name)
    c.font = F(bold=True); c.fill = PatternFill('solid', fgColor='EFEFEA')
    c.border = Border(bottom=Side(style='thin', color='D8D6CF'))
    c.alignment = Alignment(wrap_text=True, vertical='center')
for i, w in enumerate([44, 14, 14, 13, 50], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
r = 5
def row(name, a, b, fmt='#,##0.1', good_up=True, comment='', eps=0):
    global r
    dv = None if (a is None or b is None) else b - a
    vals = [name, a, b, dv, comment]
    fmts = [None, fmt, fmt, '+' + fmt + ';-' + fmt + ';0', None]
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=i, value=v); c.font = F()
        if fmts[i-1]: c.number_format = fmts[i-1]
    if dv is not None and abs(dv) > eps:
        ws.cell(row=r, column=4).font = GOOD if (dv > 0) == good_up else BAD
    r += 1

w1, w2 = H1['waterfall'], H2['waterfall']
row('Выручка бНДС, млн ₽ (по выполнению)', H1['byEnds']['s']/1e6, H2['byEnds']['s']/1e6, '#,##0.1', True, 'вторая половина без последней недели месяца')
row('Выручка, млн ₽/сутки', H1['byEnds']['s']/1e6/13, H2['byEnds']['s']/1e6/14, '#,##0.00', True, '')
row('Рейсов', H1['byEnds']['n'], H2['byEnds']['n'], '#,##0', True, '')
row('Средний чек бНДС, ₽', H1['byEnds']['s']/max(1,H1['byEnds']['n']), H2['byEnds']['s']/max(1,H2['byEnds']['n']), '#,##0', True, '')
row('На линии среднесуточно, машин', H1['avgOnline'], H2['avgOnline'], '#,##0.1', True, 'живой ряд')
row('КТГ, %', H1['park']['ktg'], H2['park']['ktg'], '#,##0.1', True, '')
row('КВЛ, %', H1['park']['kvl'], H2['park']['kvl'], '#,##0.1', True, '')
row('КИП, %', H1['park']['kip'], H2['park']['kip'], '#,##0.1', True, '')
row('Дорожная скорость, км/ч', w1['km']/w1['roadH'], w2['km']/w2['roadH'], '#,##0.1', True, 'чистые цепочки')
row('Ворота выгрузки, ч/рейс', w1['unloadH']/w1['n'], w2['unloadH']/w2['n'], '#,##0.1', False, 'цель 6 ч')
row('Ворота погрузки, ч/рейс', w1['loadH']/w1['n'], w2['loadH']/w2['n'], '#,##0.1', False, '')
row('Прибытия на погрузку вовремя, %', 100*(H1['late']['P']['n']-H1['late']['P']['late1'])/max(1,H1['late']['P']['n']), 100*(H2['late']['P']['n']-H2['late']['P']['late1'])/max(1,H2['late']['P']['n']), '#,##0', True, '')
row('Прибытия на выгрузку вовремя, %', 100*(H1['late']['D']['n']-H1['late']['D']['late1'])/max(1,H1['late']['D']['n']), 100*(H2['late']['D']['n']-H2['late']['D']['late1'])/max(1,H2['late']['D']['n']), '#,##0', True, '')
row('Следующий до выгрузки, %', H1['next']['pct'], H2['next']['pct'], '#,##0', True, 'цель 85')
row('Стык, ч (медиана)', H1['gaps']['medianH'], H2['gaps']['medianH'], '#,##0.1', False, 'цель 8')
row('Стык: ждали заказ, ч', H1['gaps']['waitOrderAvg'], H2['gaps']['waitOrderAvg'], '#,##0.1', False, '')
row('Стык: ждали окно клиента, ч', H1['gaps']['waitSlotAvg'], H2['gaps']['waitSlotAvg'], '#,##0.1', False, '')
wb.save(path)
print('OK')
